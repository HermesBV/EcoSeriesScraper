"""Copia títulos y valoraciones del inventario maestro a SeriesMacro."""

from __future__ import annotations

from pathlib import Path
import sys

from openpyxl import load_workbook


SOURCE = Path(__file__).resolve().parents[1] / "IndiceSeries.xlsx"


def sincronizar(target: Path) -> tuple[int, int]:
    source_book = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        source_sheet = source_book["Codificacion"]
        source_header = {cell.value: cell.column - 1 for cell in source_sheet[1]}
        source_rows = {}
        for row in source_sheet.iter_rows(min_row=2, values_only=True):
            key = (str(row[source_header["Código fuente"]]).strip(), str(row[source_header["ID"]]).strip())
            source_rows[key] = (row[source_header["Nombre serie"]], row[source_header["Valoración"]])
    finally:
        source_book.close()

    target_book = load_workbook(target)
    temporary = target.with_name("IndiceSeries.metadata.tmp.xlsx")
    try:
        sheet = target_book["Codificacion"]
        header = {cell.value: cell.column for cell in sheet[1]}
        target_keys = set()
        title_changes = valuation_changes = 0
        for row in sheet.iter_rows(min_row=2):
            key = (str(row[header["Código fuente"] - 1].value).strip(), str(row[header["ID"] - 1].value).strip())
            if key not in source_rows:
                raise ValueError(f"Serie de SeriesMacro ausente en el maestro: {key}")
            target_keys.add(key)
            title, valuation = source_rows[key]
            for name, value in (("Nombre serie", title), ("Valoración", valuation)):
                cell = row[header[name] - 1]
                if cell.value != value:
                    cell.value = value
                    if name == "Nombre serie":
                        title_changes += 1
                    else:
                        valuation_changes += 1
        if target_keys != set(source_rows):
            raise ValueError(f"El maestro tiene {len(set(source_rows) - target_keys)} series ajenas a SeriesMacro")
        target_book.save(temporary)
    finally:
        target_book.close()
    try:
        check = load_workbook(temporary, read_only=True)
        try:
            if check["Codificacion"].max_row != len(source_rows) + 1:
                raise ValueError("Cambió la cantidad de filas al guardar")
        finally:
            check.close()
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    return title_changes, valuation_changes


if __name__ == "__main__":
    count = sincronizar(Path(sys.argv[1]).resolve())
    print(f"Títulos actualizados: {count[0]}; valoraciones actualizadas: {count[1]}")
