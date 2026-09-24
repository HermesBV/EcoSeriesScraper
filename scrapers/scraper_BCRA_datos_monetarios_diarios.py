"""Descarga e incorpora todas las series del libro monetario diario del BCRA."""

from __future__ import annotations

from pathlib import Path
import re
import unicodedata

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


RAIZ_PROYECTO = Path(__file__).resolve().parents[1]
ARCHIVO_BD = RAIZ_PROYECTO / "BD.xlsx"
CARPETA_FUENTE = RAIZ_PROYECTO / "fuentes_BD" / "BCRA" / "DatosMonetariosDiarios"
ARCHIVO_FUENTE = CARPETA_FUENTE / "series.xlsm"
EXCEL_URL = "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/series.xlsm"
PAGINA_FUENTE = "https://www.bcra.gob.ar/publicaciones-y-estadisticas/estadisticas-datos-monetarios-diarios/"
CODIGO_FUENTE = "bcra-dmd"

HOJAS_ORIGEN = {
    "BASE MONETARIA": "BCRA Base Monetaria D",
    "RESERVAS": "BCRA Reservas D",
    "DEPOSITOS": "BCRA Depositos D",
    "PRESTAMOS": "BCRA Prestamos D",
    "TASAS DE MERCADO": "BCRA Tasas Mercado D",
    "INSTRUMENTOS DEL BCRA": "BCRA Instrumentos D",
}
HOJAS_ADMINISTRADAS = set(HOJAS_ORIGEN.values())
ALIAS_HOJAS_API = {
    "INTRUMENTOS DEL BCRA": "INSTRUMENTOS DEL BCRA",
    "INSTRUMENTOS": "INSTRUMENTOS DEL BCRA",
}


def _sesion() -> requests.Session:
    sesion = requests.Session()
    sesion.headers["User-Agent"] = "EcoSeriesScraper/1.0 (datos publicos BCRA)"
    return sesion


def _excel_valido(ruta: Path) -> bool:
    if not ruta.is_file() or ruta.stat().st_size < 100_000:
        return False
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True, keep_vba=True)
        valido = set(HOJAS_ORIGEN) | {"API_Series"} <= set(libro.sheetnames)
        libro.close()
        return valido
    except Exception:
        return False


def descargar_excel() -> Path:
    """Publica la descarga atomica y reutiliza una copia local valida ante fallos."""
    CARPETA_FUENTE.mkdir(parents=True, exist_ok=True)
    temporal = ARCHIVO_FUENTE.with_name("series.descarga.xlsm")
    try:
        with _sesion().get(EXCEL_URL, stream=True, timeout=(20, 240)) as respuesta:
            respuesta.raise_for_status()
            with temporal.open("wb") as archivo:
                for bloque in respuesta.iter_content(64 * 1024):
                    if bloque:
                        archivo.write(bloque)
        if not _excel_valido(temporal):
            raise ValueError("La descarga no es el libro de datos monetarios diarios del BCRA")
        temporal.replace(ARCHIVO_FUENTE)
    except (OSError, requests.RequestException, ValueError):
        if not _excel_valido(ARCHIVO_FUENTE):
            raise
        print("BCRA DMD: se reutiliza la copia local de series.xlsm", flush=True)
    finally:
        temporal.unlink(missing_ok=True)
    return ARCHIVO_FUENTE


def _texto(valor: object) -> str:
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def _slug(valor: str) -> str:
    normal = unicodedata.normalize("NFKD", valor).encode("ascii", "ignore").decode().casefold()
    return re.sub(r"[^a-z0-9]+", "-", normal).strip("-")


def _unidad(nombre: str) -> str:
    coincidencia = re.search(r"\((en [^)]+)\)\s*$", nombre, flags=re.IGNORECASE)
    if coincidencia:
        return coincidencia.group(1)[3:].strip()
    texto = nombre.casefold()
    if "tipo de cambio" in texto:
        return "Pesos por USD"
    return "No informado"


def leer_catalogo_api(ruta: Path) -> pd.DataFrame:
    """Lee los IDs que el propio BCRA vincula con las columnas del libro."""
    catalogo = pd.read_excel(ruta, sheet_name="API_Series", header=3, usecols="A:D")
    catalogo.columns = ["id", "nombre", "hoja", "columna"]
    catalogo["id"] = pd.to_numeric(catalogo["id"], errors="coerce")
    catalogo = catalogo.dropna(subset=["id", "nombre", "hoja", "columna"]).copy()
    catalogo["id"] = catalogo["id"].astype(int).astype(str)
    catalogo["nombre"] = catalogo["nombre"].map(_texto)
    catalogo["hoja"] = catalogo["hoja"].map(_texto).replace(ALIAS_HOJAS_API)
    catalogo["columna"] = catalogo["columna"].map(lambda valor: _texto(valor).upper())
    catalogo = catalogo[catalogo["hoja"].isin(HOJAS_ORIGEN)]
    if catalogo["id"].duplicated().any() or catalogo[["hoja", "columna"]].duplicated().any():
        raise ValueError("API_Series contiene IDs o columnas duplicadas")
    return catalogo.reset_index(drop=True)


def _encabezado_estable(hoja, columna: int) -> str:
    partes = [_texto(hoja.cell(fila, columna).value) for fila in range(4, 8)]
    partes = [parte for indice, parte in enumerate(partes) if parte and parte not in partes[:indice]]
    return " - ".join(partes) or _texto(hoja.cell(9, columna).value) or get_column_letter(columna)


def _columnas_publicadas(hoja, catalogo_hoja: pd.DataFrame) -> list[dict[str, str]]:
    por_columna = {fila.columna: fila for fila in catalogo_hoja.itertuples(index=False)}
    columnas: list[dict[str, str]] = []
    for numero in range(2, hoja.max_column + 1):
        letra = get_column_letter(numero)
        api = por_columna.get(letra)
        tiene_datos = any(
            isinstance(hoja.cell(fila, numero).value, (int, float))
            for fila in range(10, hoja.max_row + 1)
        )
        if api is None and not tiene_datos:
            continue
        nombre = api.nombre if api is not None else _encabezado_estable(hoja, numero)
        identificador = api.id if api is not None else f"{_slug(hoja.title)}-{letra.casefold()}"
        columnas.append({"letra": letra, "id": identificador, "nombre": nombre})
    return columnas


def construir_salida(ruta: Path) -> tuple[dict[str, pd.DataFrame], pd.DataFrame]:
    catalogo = leer_catalogo_api(ruta)
    # El acceso por columnas y por encabezados es mucho mas rapido con la hoja
    # materializada; el libro oficial ronda pocos MB y entra holgadamente en memoria.
    libro = load_workbook(ruta, read_only=False, data_only=True, keep_vba=True)
    hojas_salida: dict[str, pd.DataFrame] = {}
    inventario: list[dict[str, object]] = []
    try:
        for hoja_origen, hoja_bd in HOJAS_ORIGEN.items():
            hoja = libro[hoja_origen]
            columnas = _columnas_publicadas(hoja, catalogo[catalogo["hoja"].eq(hoja_origen)])
            indices = [1] + [column_index_from_string(item["letra"]) for item in columnas]
            filas = ([hoja.cell(fila, columna).value for columna in indices]
                     for fila in range(10, hoja.max_row + 1))
            datos = pd.DataFrame(filas, columns=["fecha"] + [item["nombre"] for item in columnas])
            datos["fecha"] = pd.to_datetime(datos["fecha"], errors="coerce")
            for columna in datos.columns[1:]:
                datos[columna] = pd.to_numeric(datos[columna], errors="coerce")
            datos = datos.dropna(subset=["fecha"]).dropna(subset=datos.columns[1:], how="all")
            datos = datos.sort_values("fecha").drop_duplicates("fecha", keep="last").reset_index(drop=True)
            if datos.empty:
                raise ValueError(f"{hoja_origen} no contiene observaciones")
            hojas_salida[hoja_bd] = datos
            for item in columnas:
                serie = datos[item["nombre"]]
                fechas = datos.loc[serie.notna(), "fecha"]
                inventario.append({
                    "ID": item["id"], "Código fuente": CODIGO_FUENTE,
                    "Nombre serie": item["nombre"], "Variable": _slug(item["nombre"]),
                    "Unidades": _unidad(item["nombre"]), "Valoración": "No aplica / no informado",
                    "Descripción": item["nombre"], "Frecuencia": "D",
                    "Pestaña BD": hoja_bd, "Columna BD": item["nombre"],
                    "Archivo origen": ARCHIVO_FUENTE.name, "Hoja origen": hoja_origen,
                    "Origen": "Banco Central de la República Argentina (BCRA)",
                    "Fuente": PAGINA_FUENTE, "Catálogo ID": "bcra-api-series",
                    "Dataset ID": "datos-monetarios-diarios", "Título dataset": "Datos monetarios diarios",
                    "Tema dataset": f"BCRA / Estadísticas monetarias y financieras / {hoja_origen.title()}",
                    "Responsable dataset": "Banco Central de la República Argentina (BCRA)",
                    "Fuente de valores": "Excel BCRA", "Fecha inicio": fechas.min() if len(fechas) else None,
                    "Fecha fin": fechas.max() if len(fechas) else None, "Estado": "VIGENTE",
                })
    finally:
        libro.close()
    inventario_df = pd.DataFrame(inventario)
    if inventario_df[["Código fuente", "ID"]].duplicated().any():
        raise ValueError("Se generaron identidades duplicadas")
    return hojas_salida, inventario_df


def procesar(descargar: bool = True) -> dict[str, int]:
    ruta = descargar_excel() if descargar else ARCHIVO_FUENTE
    if not _excel_valido(ruta):
        raise FileNotFoundError(f"No existe un Excel BCRA válido en {ruta}")
    hojas, inventario = construir_salida(ruta)
    from scrapers.scraper_IED import guardar_datos_preservando_formato
    guardar_datos_preservando_formato(
        ARCHIVO_BD, hojas, frecuencias={nombre: "D" for nombre in hojas},
        hojas_obsoletas=HOJAS_ADMINISTRADAS - set(hojas),
    )
    from tools.generar_codificacion import generar
    generar(inventario)
    resumen = {"series": len(inventario), "hojas": len(hojas)}
    print(f"BCRA datos monetarios diarios terminado: {resumen}", flush=True)
    return resumen


def ejecutar() -> None:
    procesar(descargar=True)


if __name__ == "__main__":
    ejecutar()
