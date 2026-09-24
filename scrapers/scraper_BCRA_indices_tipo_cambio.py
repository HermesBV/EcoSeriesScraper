"""Descarga e incorpora los índices de tipo de cambio publicados por el BCRA."""

from __future__ import annotations

from pathlib import Path
import re
import unicodedata

import pandas as pd
import requests
from openpyxl import load_workbook


RAIZ_PROYECTO = Path(__file__).resolve().parents[1]
ARCHIVO_BD = RAIZ_PROYECTO / "BD.xlsx"
CARPETA_FUENTE = RAIZ_PROYECTO / "fuentes_BD" / "BCRA" / "IndicesTipoCambio"
PAGINA_FUENTE = "https://www.bcra.gob.ar/indices-de-tipo-de-cambio-multilateral/"
CODIGO_FUENTE = "bcra-itc"
EXCEL_URLS = {
    "ITCRMSerie.xlsx": "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/ITCRMSerie.xlsx",
    "ITCNMSerie.xlsx": "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/ITCNMSerie.xlsx",
}
HOJAS = {
    ("ITCRMSerie.xlsx", "ITCRM y bilaterales"): ("BCRA ITCRM D", "D", "Índices de tipo de cambio real"),
    ("ITCRMSerie.xlsx", "ITCRM y bilaterales prom. mens."): ("BCRA ITCRM M", "M", "Índices de tipo de cambio real"),
    ("ITCNMSerie.xlsx", "ITCNM y bilaterales"): ("BCRA ITCNM D", "D", "Índices de tipo de cambio nominal"),
    ("ITCNMSerie.xlsx", "ITCNM y bilaterales prom. mens."): ("BCRA ITCNM M", "M", "Índices de tipo de cambio nominal"),
    ("ITCRMSerie.xlsx", "Ponderadores"): ("BCRA Ponderadores M", "M", "Ponderadores de comercio de manufacturas"),
}
HOJAS_ADMINISTRADAS = {config[0] for config in HOJAS.values()}


def _sesion() -> requests.Session:
    sesion = requests.Session()
    sesion.headers["User-Agent"] = "EcoSeriesScraper/1.0 (datos públicos BCRA)"
    return sesion


def _excel_valido(ruta: Path) -> bool:
    if not ruta.is_file() or ruta.stat().st_size < 1_000:
        return False
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
        valido = all(hoja in libro.sheetnames for archivo, hoja in HOJAS if archivo == ruta.name)
        libro.close()
        return valido
    except Exception:
        return False


def descargar_excels() -> dict[str, Path]:
    """Publica cada descarga atómicamente y reutiliza la última copia si la red falla."""
    CARPETA_FUENTE.mkdir(parents=True, exist_ok=True)
    resultados: dict[str, Path] = {}
    with _sesion() as sesion:
        for nombre, url in EXCEL_URLS.items():
            destino = CARPETA_FUENTE / nombre
            temporal = destino.with_name(f"{destino.stem}.descarga.xlsx")
            try:
                with sesion.get(url, stream=True, timeout=(20, 180)) as respuesta:
                    respuesta.raise_for_status()
                    with temporal.open("wb") as archivo:
                        for bloque in respuesta.iter_content(64 * 1024):
                            if bloque:
                                archivo.write(bloque)
                if not _excel_valido(temporal):
                    raise ValueError(f"La descarga de {nombre} no es un Excel BCRA válido")
                temporal.replace(destino)
            except (OSError, requests.RequestException, ValueError):
                if not _excel_valido(destino):
                    raise
                print(f"BCRA ITC: se reutiliza la copia local de {nombre}", flush=True)
            finally:
                temporal.unlink(missing_ok=True)
            resultados[nombre] = destino
    return resultados


def _texto(valor: object) -> str:
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def _slug(valor: str) -> str:
    normal = unicodedata.normalize("NFKD", valor).encode("ascii", "ignore").decode().casefold()
    return re.sub(r"[^a-z0-9]+", "-", normal).strip("-")


def leer_hoja(ruta: Path, hoja: str, frecuencia: str) -> pd.DataFrame:
    """Lee el bloque tabular, elimina adornos y normaliza fechas al inicio del período."""
    datos = pd.read_excel(ruta, sheet_name=hoja, header=1)
    datos = datos.dropna(axis=1, how="all").dropna(how="all")
    if datos.empty or len(datos.columns) < 2:
        raise ValueError(f"{ruta.name}/{hoja} no contiene series")
    columnas = [_texto(columna) for columna in datos.columns]
    columnas[0] = "fecha"
    datos.columns = columnas
    datos["fecha"] = pd.to_datetime(datos["fecha"], errors="coerce")
    for columna in columnas[1:]:
        datos[columna] = pd.to_numeric(datos[columna], errors="coerce")
    datos = datos.dropna(subset=["fecha"])
    datos = datos.dropna(subset=columnas[1:], how="all")
    if frecuencia == "M":
        datos["fecha"] = datos["fecha"].dt.to_period("M").dt.to_timestamp()
    if datos["fecha"].duplicated().any() or not datos["fecha"].is_monotonic_increasing:
        raise ValueError(f"Fechas duplicadas o desordenadas en {ruta.name}/{hoja}")
    return datos.reset_index(drop=True)


def construir_salida(archivos: dict[str, Path]) -> tuple[dict[str, pd.DataFrame], pd.DataFrame]:
    hojas_salida: dict[str, pd.DataFrame] = {}
    inventario: list[dict[str, object]] = []
    for (archivo, hoja_origen), (hoja_bd, frecuencia, tema) in HOJAS.items():
        datos = leer_hoja(archivos[archivo], hoja_origen, frecuencia)
        hojas_salida[hoja_bd] = datos
        es_ponderador = hoja_origen == "Ponderadores"
        for columna in datos.columns[1:]:
            id_origen = f"{archivo}|{hoja_origen}|{columna}"
            nombre = f"{columna} ({'mensual' if frecuencia == 'M' else 'diario'})"
            descripcion = (
                "Participación porcentual del socio comercial en el promedio móvil de 12 meses "
                "del comercio argentino de manufacturas."
                if es_ponderador else
                f"{columna}, índice Laspeyres geométrico encadenado publicado por el BCRA."
            )
            inventario.append({
                "ID": id_origen,
                "Código fuente": CODIGO_FUENTE,
                "Nombre serie": nombre,
                "Variable": _slug(columna),
                "Unidades": "Porcentaje" if es_ponderador else "Índice 17-dic-2015=100",
                "Valoración": "No aplica / no informado",
                "Descripción": descripcion,
                "Frecuencia": frecuencia,
                "Pestaña BD": hoja_bd,
                "Columna BD": columna,
                "Archivo origen": archivo,
                "Hoja origen": hoja_origen,
                "Origen": "Banco Central de la República Argentina (BCRA)",
                "Fuente": PAGINA_FUENTE,
                "Catálogo ID": "bcra-indices-tipo-cambio",
                "Dataset ID": "ITCRM" if "ITCRM" in archivo else "ITCNM",
                "Título dataset": tema,
                "Tema dataset": "Tipo de cambio",
                "Responsable dataset": "Banco Central de la República Argentina (BCRA)",
                "Fuente de valores": "Excel BCRA",
                "Fecha inicio": datos.loc[datos[columna].notna(), "fecha"].min(),
                "Fecha fin": datos.loc[datos[columna].notna(), "fecha"].max(),
                "Estado": "VIGENTE",
            })
    return hojas_salida, pd.DataFrame(inventario)


def _guardar_hojas(hojas: dict[str, pd.DataFrame]) -> None:
    from scrapers.scraper_IED import guardar_datos_preservando_formato
    frecuencias = {nombre: ("D" if nombre.endswith(" D") else "M") for nombre in hojas}
    guardar_datos_preservando_formato(
        ARCHIVO_BD, hojas, frecuencias=frecuencias,
        hojas_obsoletas=HOJAS_ADMINISTRADAS - set(hojas),
    )


def procesar(descargar: bool = True) -> dict[str, int]:
    archivos = descargar_excels() if descargar else {
        nombre: CARPETA_FUENTE / nombre for nombre in EXCEL_URLS
    }
    faltantes = [nombre for nombre, ruta in archivos.items() if not _excel_valido(ruta)]
    if faltantes:
        raise FileNotFoundError(f"Faltan Excel BCRA válidos: {', '.join(faltantes)}")
    hojas, inventario = construir_salida(archivos)
    _guardar_hojas(hojas)
    from tools.generar_codificacion import generar
    from tools.generar_codificacion import cargar_indice
    actual = cargar_indice()
    actual = actual[actual["Código fuente"].astype(str).ne(CODIGO_FUENTE)]
    generar(pd.concat([actual, inventario], ignore_index=True))
    resumen = {"series": len(inventario), "hojas": len(hojas)}
    print(f"BCRA índices de tipo de cambio terminado: {resumen}", flush=True)
    return resumen


def ejecutar() -> None:
    procesar(descargar=True)


if __name__ == "__main__":
    ejecutar()
